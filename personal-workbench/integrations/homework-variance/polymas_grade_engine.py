#!/usr/bin/env python3
"""
Config-driven Polymas 作业批阅引擎。

用法:
  python polymas_grade_engine.py init --config job.json
  python polymas_grade_engine.py next --config job.json
  python polymas_grade_engine.py status --config job.json
  python polymas_grade_engine.py report --config job.json
  python polymas_grade_engine.py run --config job.json   # 循环 next 直到 done
"""
from __future__ import annotations

import argparse
import json
import mimetypes
import os
import re
import statistics
import sys
import time
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

# Windows 控制台/管道默认常为 GBK；强制 stdout/stderr 用 UTF-8，避免中文 level 名乱码
def _force_utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
        except Exception:
            pass


_force_utf8_stdio()

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HOST = "https://cloudapi.polymas.com"
SUCCESS = {"finished", "completed", "success", "done", "FINISHED", "COMPLETED", "SUCCESS", "DONE"}
FAIL = {"failed", "fail", "error", "canceled", "rejected", "FAILED", "FAIL", "ERROR"}

# 平台 homeworkFileAnalysis 常见空答案标记
_EMPTY_MARKERS = (
    "未获取到作答内容",
    "未获取到",
    "未提取到",
    "未识别到",
    "无作答内容",
    "无作答",
    "无答案",
    "暂无答案",
    "暂无",
    "空",
    "none",
    "null",
    "n/a",
)

DEFAULT_LLM_BASE = "https://llm-service.polymas.com/api/openai/v1"
DEFAULT_LLM_MODEL = "claude-sonnet-4-6"


def load_cfg(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def save_cfg(path: Path, cfg: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(cfg, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def state_path(cfg: dict) -> Path:
    return Path(cfg["out_dir"]) / "state.json"


def runs_dir(cfg: dict) -> Path:
    d = Path(cfg["out_dir"]) / "runs"
    d.mkdir(parents=True, exist_ok=True)
    return d


def load_state(cfg: dict) -> dict:
    return json.loads(state_path(cfg).read_text(encoding="utf-8"))


def save_state(cfg: dict, st: dict) -> None:
    p = state_path(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(st, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def sess(cfg: dict) -> requests.Session:
    s = requests.Session()
    auth = cfg["auth"]["jwt"]
    cookie = cfg["auth"].get("cookie") or f"ai-poly={auth}"
    if "ai-poly=" not in cookie:
        cookie = f"ai-poly={auth}; {cookie}"
    s.headers.update(
        {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Accept": "application/json, text/plain, */*",
            "Authorization": auth,
            "ai-poly": auth,
            "Cookie": cookie,
            "Origin": "https://hike-teaching-center.polymas.com",
            "Referer": "https://hike-teaching-center.polymas.com/",
        }
    )
    return s


def get_version(s: requests.Session, instance_nid: str, agent_id: str):
    r = s.post(
        f"{HOST}/agents/v1/agent/details",
        json={"instanceIds": [instance_nid], "needToToolSchema": False, "version": ""},
        timeout=60,
    )
    data = r.json()
    details = ((data.get("data") or {}).get("instanceDetails") or [])
    if not details:
        raise RuntimeError(f"agent/details failed: {data}")
    return details[0].get("version"), details[0].get("agentNid") or agent_id


def extract_scores(task_payload: dict) -> dict:
    d = task_payload.get("data") or {}
    arts = d.get("artifacts") or []
    result = {"total": None, "items": {}, "overall_comment": None, "raw_names": [], "max_items": {}}
    if not arts:
        return result
    try:
        data_part = arts[0]["parts"][0]["data"]
    except Exception:
        data_part = arts[0]
    if isinstance(data_part, str):
        try:
            data_part = json.loads(data_part)
        except Exception:
            return result
    if not isinstance(data_part, dict):
        return result
    for k in ("totalScore", "total_score", "finalScore"):
        if isinstance(data_part.get(k), (int, float)):
            result["total"] = float(data_part[k])
            break
    result["overall_comment"] = data_part.get("overallComment")
    qs = data_part.get("questionScores") or []
    if isinstance(qs, list):
        for i, it in enumerate(qs):
            if not isinstance(it, dict):
                continue
            name = it.get("name") or f"Q{i+1}"
            result["raw_names"].append(name)
            if isinstance(it.get("score"), (int, float)):
                result["items"][str(name)] = float(it["score"])
            if isinstance(it.get("totalScore"), (int, float)):
                result["max_items"][str(name)] = float(it["totalScore"])
    if result["total"] is None and result["items"]:
        result["total"] = float(sum(result["items"].values()))
    return result


def build_text_input(parsed) -> str | None:
    content = parsed.get("content") if isinstance(parsed, dict) else None
    if not isinstance(content, list):
        if isinstance(parsed, dict) and isinstance(parsed.get("text"), str):
            return parsed["text"]
        return None
    parts = []
    for item in content:
        name = item.get("itemName") or item.get("scoreItemName") or ""
        ans = item.get("stuAnswerContent") or ""
        parts.append(f"【{name}】\n{ans}")
    return "\n\n".join(parts) if parts else None


def is_empty_answer(text) -> bool:
    if text is None:
        return True
    s = str(text).strip()
    if not s:
        return True
    low = s.lower()
    for m in _EMPTY_MARKERS:
        if s == m or low == m.lower() or s.startswith(m) or low.startswith(m.lower()):
            return True
    # 极短且无实质内容
    if len(s) <= 2 and s in {".", "-", "—", "/", "无", "空"}:
        return True
    return False


def content_items(parsed) -> list[dict]:
    if not isinstance(parsed, dict):
        return []
    c = parsed.get("content")
    return c if isinstance(c, list) else []


def detect_empty_items(parsed) -> list[dict]:
    """返回空答案题列表: [{index, name, old}]"""
    out = []
    for i, item in enumerate(content_items(parsed)):
        if not isinstance(item, dict):
            continue
        name = item.get("itemName") or item.get("scoreItemName") or f"Q{i+1}"
        ans = item.get("stuAnswerContent")
        if is_empty_answer(ans):
            out.append({"index": i, "name": str(name), "old": "" if ans is None else str(ans)})
    return out


def extract_local_file_text(path: Path, max_chars: int = 48000) -> str:
    """从本地答卷提取纯文本，供 LLM 回填空题（不依赖额外包）。"""
    if not path or not Path(path).is_file():
        return ""
    fp = Path(path)
    suf = fp.suffix.lower()
    try:
        if suf in {".txt", ".md", ".csv"}:
            return fp.read_text(encoding="utf-8", errors="replace")[:max_chars]
        if suf in {".docx"}:
            return _extract_docx_text(fp)[:max_chars]
        if suf in {".doc", ".pdf", ".ppt", ".pptx", ".png", ".jpg", ".jpeg"}:
            # 二进制格式：至少给出文件名提示，LLM 仍可能无法恢复
            return f"[二进制文件 {fp.name}，未能本地抽取正文；请仅根据已解析到的非空题上下文推断空题，若无法确定则保留空]"
    except Exception as e:
        return f"[抽取失败: {e}]"
    return ""


def _extract_docx_text(fp: Path) -> str:
    """粗提取 docx 正文（paragraph + table cell）。"""
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    chunks: list[str] = []
    with zipfile.ZipFile(fp) as zf:
        # 主文档
        for name in ("word/document.xml",):
            if name not in zf.namelist():
                continue
            root = ET.fromstring(zf.read(name))
            for p in root.iter(f"{{{ns['w']}}}p"):
                texts = [t.text or "" for t in p.iter(f"{{{ns['w']}}}t")]
                line = "".join(texts).strip()
                if line:
                    chunks.append(line)
        # 页眉页脚有时含题号，一并纳入
        for name in zf.namelist():
            if not name.startswith("word/header") and not name.startswith("word/footer"):
                continue
            if not name.endswith(".xml"):
                continue
            try:
                root = ET.fromstring(zf.read(name))
                for p in root.iter(f"{{{ns['w']}}}p"):
                    texts = [t.text or "" for t in p.iter(f"{{{ns['w']}}}t")]
                    line = "".join(texts).strip()
                    if line:
                        chunks.append(line)
            except Exception:
                pass
    # 去连续重复行
    out: list[str] = []
    prev = None
    for c in chunks:
        if c != prev:
            out.append(c)
        prev = c
    return "\n".join(out)


def _engine_root() -> Path:
    return Path(__file__).resolve().parent


def load_secrets_file() -> dict:
    candidates = []
    configured = os.environ.get("PERSONAL_WORKBENCH_HOMEWORK_VARIANCE_SECRETS", "").strip()
    if configured:
        candidates.append(Path(configured).expanduser())
    candidates.append(_engine_root() / "secrets.json")
    for p in candidates:
        if not p.is_file():
            continue
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except Exception:
            continue
    return {}


def resolve_llm_cfg(cfg: dict | None = None) -> dict | None:
    """
    解析 LLM 配置。优先级：
      job.json llm.*  → 环境变量 POLY_LLM_*  → secrets.json llm.*
    返回 {api_key, base_url, model} 或 None。
    """
    cfg = cfg or {}
    llm = cfg.get("llm") if isinstance(cfg.get("llm"), dict) else {}
    secrets = load_secrets_file()
    sec_llm = secrets.get("llm") if isinstance(secrets.get("llm"), dict) else {}

    api_key = (
        (llm.get("api_key") or "").strip()
        or os.environ.get("POLY_LLM_API_KEY", "").strip()
        or (sec_llm.get("api_key") or "").strip()
        or (secrets.get("llm_api_key") or "").strip()
    )
    if not api_key or api_key.startswith("粘贴") or "your-" in api_key.lower():
        return None

    base = (
        (llm.get("base_url") or "").strip()
        or os.environ.get("POLY_LLM_BASE_URL", "").strip()
        or (sec_llm.get("base_url") or "").strip()
        or DEFAULT_LLM_BASE
    ).rstrip("/")
    # 允许用户写 .../v1 或 .../v1/chat/completions
    if base.endswith("/chat/completions"):
        base = base[: -len("/chat/completions")].rstrip("/")

    model = (
        (llm.get("model") or "").strip()
        or os.environ.get("POLY_LLM_MODEL", "").strip()
        or (sec_llm.get("model") or "").strip()
        or DEFAULT_LLM_MODEL
    )
    return {"api_key": api_key, "base_url": base, "model": model}


def llm_chat(llm_cfg: dict, messages: list[dict], temperature: float = 0.1, max_tokens: int = 8000) -> str:
    url = f"{llm_cfg['base_url']}/chat/completions"
    headers = {
        "Authorization": f"Bearer {llm_cfg['api_key']}",
        "Content-Type": "application/json",
    }
    body = {
        "model": llm_cfg["model"],
        "messages": messages,
        "temperature": temperature,
        "max_tokens": max_tokens,
    }
    r = requests.post(url, headers=headers, json=body, timeout=180)
    if r.status_code >= 400:
        raise RuntimeError(f"LLM HTTP {r.status_code}: {r.text[:500]}")
    data = r.json()
    try:
        return data["choices"][0]["message"]["content"]
    except Exception as e:
        raise RuntimeError(f"LLM bad response: {e}; body={str(data)[:400]}") from e


def _parse_llm_json_object(text: str) -> dict:
    """从模型输出中抠出 JSON 对象。"""
    s = (text or "").strip()
    if not s:
        return {}
    # 去掉 ```json 围栏
    if "```" in s:
        m = re.search(r"```(?:json)?\s*([\s\S]*?)```", s, flags=re.I)
        if m:
            s = m.group(1).strip()
    try:
        data = json.loads(s)
        return data if isinstance(data, dict) else {}
    except Exception:
        pass
    # 找最外层 {}
    i = s.find("{")
    j = s.rfind("}")
    if i >= 0 and j > i:
        try:
            data = json.loads(s[i : j + 1])
            return data if isinstance(data, dict) else {}
        except Exception:
            pass
    return {}


def fix_empty_answers_with_llm(
    parsed: dict,
    file_path: str | Path,
    llm_cfg: dict,
    empty_items: list[dict] | None = None,
) -> tuple[dict, list[dict]]:
    """
    用本地文件正文 + LLM，回填 analysis 中的空答案。
    返回 (fixed_parsed, fixes[{name, old, new}])。
    """
    import copy

    empty_items = empty_items if empty_items is not None else detect_empty_items(parsed)
    if not empty_items:
        return parsed, []

    local_text = extract_local_file_text(Path(file_path))
    # 已有非空题作上下文
    known = []
    for item in content_items(parsed):
        if not isinstance(item, dict):
            continue
        name = item.get("itemName") or item.get("scoreItemName") or ""
        ans = item.get("stuAnswerContent") or ""
        if not is_empty_answer(ans):
            known.append({"name": str(name), "answer": str(ans)[:800]})

    empty_names = [e["name"] for e in empty_items]
    system = (
        "你是作业答卷结构化抽取助手。平台解析器漏抽了部分题目的学生答案。"
        "请根据【答卷原文】把漏掉的题号对应答案补全。"
        "规则：\n"
        "1. 只输出一个 JSON 对象，不要 markdown，不要解释。\n"
        "2. key 必须是题号/题目名（与给定列表完全一致），value 是学生答案原文（可含公式文字描述）。\n"
        "3. 严禁编造学生没写的内容；原文确实没有则 value 用空字符串 \"\"。\n"
        "4. 不要批改、不要给分、不要改写已有非空题。"
    )
    user = {
        "empty_items": empty_names,
        "known_non_empty": known[:40],
        "source_file": Path(file_path).name,
        "document_text": local_text[:45000] if local_text else "",
    }
    content = llm_chat(
        llm_cfg,
        [
            {"role": "system", "content": system},
            {
                "role": "user",
                "content": "请补全下列空题答案，只返回 JSON：\n"
                + json.dumps(user, ensure_ascii=False),
            },
        ],
        temperature=0.1,
        max_tokens=8000,
    )
    mapping = _parse_llm_json_object(content)
    # 接受 {"fixes":[{name,answer}]} 或 {题号: 答案} 两种形态
    if isinstance(mapping.get("fixes"), list):
        tmp = {}
        for it in mapping["fixes"]:
            if isinstance(it, dict):
                n = it.get("name") or it.get("itemName") or it.get("key")
                a = it.get("answer") or it.get("stuAnswerContent") or it.get("value")
                if n is not None:
                    tmp[str(n)] = "" if a is None else str(a)
        mapping = tmp
    # 规范化 key
    norm_map: dict[str, str] = {}
    for k, v in mapping.items():
        if isinstance(v, (dict, list)):
            continue
        norm_map[str(k).strip()] = "" if v is None else str(v)

    fixed = copy.deepcopy(parsed)
    fixes: list[dict] = []
    for e in empty_items:
        name = e["name"]
        new_ans = None
        if name in norm_map:
            new_ans = norm_map[name]
        else:
            # 宽松匹配：去空格 / 题前缀
            for k, v in norm_map.items():
                if k.replace(" ", "") == name.replace(" ", ""):
                    new_ans = v
                    break
                if name.endswith(k) or k.endswith(name):
                    new_ans = v
                    break
        if new_ans is None:
            continue
        new_ans = new_ans.strip()
        if not new_ans or is_empty_answer(new_ans):
            continue
        idx = e["index"]
        items = content_items(fixed)
        if 0 <= idx < len(items) and isinstance(items[idx], dict):
            old = items[idx].get("stuAnswerContent")
            items[idx]["stuAnswerContent"] = new_ans
            fixes.append(
                {
                    "name": name,
                    "old": "" if old is None else str(old)[:80],
                    "new": new_ans[:200] + ("…" if len(new_ans) > 200 else ""),
                    "new_len": len(new_ans),
                }
            )
    return fixed, fixes


def analysis_dir(cfg: dict) -> Path:
    d = Path(cfg["out_dir"]) / "analysis"
    d.mkdir(parents=True, exist_ok=True)
    return d


def cmd_init(cfg_path: Path) -> None:
    cfg = load_cfg(cfg_path)
    out = Path(cfg["out_dir"])
    out.mkdir(parents=True, exist_ok=True)
    runs_dir(cfg)
    s = sess(cfg)
    version, agent_id = get_version(s, cfg["instance_nid"], cfg.get("agent_id") or "zsxGPgjvWx")
    times = int(cfg.get("times") or 3)
    jobs = []
    for fi, f in enumerate(cfg["files"]):
        label = f.get("label") or Path(f["path"]).stem
        jobs.append(
            {
                "file_index": fi,
                "level": label,
                "filename": Path(f["path"]).name,
                "path": f["path"],
                "phase": "upload",
                "file_data": None,
                "text_input": None,
                "current_run": 1,
                "runs": [
                    {
                        "run_index": i,
                        "phase": "pending",
                        "task_id": None,
                        "total_score": None,
                        "items": {},
                        "max_items": {},
                        "state": None,
                        "timings": {},
                        "poll_count": 0,
                        "poll_t0": None,
                        "error": None,
                    }
                    for i in range(1, times + 1)
                ],
            }
        )
    st = {
        "version": version,
        "agent_id": agent_id,
        "instance_nid": cfg["instance_nid"],
        "user_nid": cfg["user_nid"],
        "times": times,
        "poll_interval": float(cfg.get("poll_interval") or 3.0),
        "job_index": 0,
        "jobs": jobs,
        "wall_start": time.time(),
        "done": False,
        "title": cfg.get("title") or "",
        "course_id": cfg.get("course_id") or "",
    }
    save_state(cfg, st)
    print(json.dumps({"ok": True, "files": len(jobs), "times": times, "version": version, "out": str(out)}, ensure_ascii=False))


def cur_job(st):
    return st["jobs"][st["job_index"]]


def cur_run(job):
    return job["runs"][job["current_run"] - 1]


def _advance(st, job, times: int):
    ri = job["current_run"]
    if ri < times:
        job["current_run"] = ri + 1
        job["phase"] = "execute"
    else:
        job["phase"] = "done"
        if st["job_index"] + 1 < len(st["jobs"]):
            st["job_index"] += 1
        else:
            st["done"] = True
            st["wall_end"] = time.time()
            st["wall_seconds"] = round(st["wall_end"] - st["wall_start"], 2)


def cmd_next(cfg_path: Path, max_polls: int = 5, budget_s: float = 38.0) -> int:
    cfg = load_cfg(cfg_path)
    st = load_state(cfg)
    if st.get("done"):
        print(json.dumps({"done": True}, ensure_ascii=False))
        return 0

    s = sess(cfg)
    deadline = time.time() + budget_s
    steps = 0
    last_code = 0
    times = int(st.get("times") or 3)
    interval = float(st.get("poll_interval") or 3.0)
    instance = st["instance_nid"]
    user_nid = st["user_nid"]

    while time.time() < deadline and steps < 12 and not st.get("done"):
        job = cur_job(st)
        phase = job["phase"]
        level = job["level"]
        ri = job["current_run"]
        print(f"[{level} run{ri}/{times}] phase={phase}", flush=True)
        steps += 1

        if phase == "upload":
            t0 = time.time()
            fp = Path(job["path"])
            if not fp.is_file():
                raise RuntimeError(f"missing file: {fp}")
            mime = mimetypes.guess_type(str(fp))[0] or "application/octet-stream"
            with open(fp, "rb") as f:
                r = s.post(
                    f"{HOST}/basic-resource/file/upload?hidden=false",
                    files={"file": (fp.name, f, mime)},
                    timeout=180,
                )
            data = r.json()
            if data.get("code") not in (200, "200"):
                raise RuntimeError(f"upload: {data}")
            job["file_data"] = data["data"]
            job["upload_s"] = round(time.time() - t0, 2)
            job["phase"] = "analyze"
            print(f"  upload {job['upload_s']}s fileId={job['file_data'].get('fileId')}", flush=True)
            save_state(cfg, st)
            continue

        if phase == "analyze":
            t0 = time.time()
            fd = job["file_data"]
            body = {
                "agentId": st["agent_id"],
                "instanceNid": instance,
                "userNid": user_nid,
                "activeMode": 1,
                "writingRequirement": "",
                "editorContent": "",
                "version": st["version"],
                "fileList": [
                    {
                        "fileName": fd.get("fileName"),
                        "fileUrl": fd.get("ossUrl") or fd.get("url"),
                        "fileId": fd.get("fileId"),
                        "ossUrl": fd.get("ossUrl"),
                    }
                ],
            }
            r = s.post(f"{HOST}/agents/v1/file/homeworkFileAnalysis", json=body, timeout=300)
            data = r.json()
            if data.get("code") not in (200, "200"):
                raise RuntimeError(f"analyze: {data}")
            raw = data.get("data")
            if isinstance(raw, str):
                try:
                    parsed = json.loads(raw)
                except Exception:
                    parsed = {"content": raw}
            else:
                parsed = raw if raw is not None else {}

            # 落盘原始解析
            adir = analysis_dir(cfg)
            safe = re.sub(r"[^\w一-鿿\-]+", "_", level)[:80] or f"file{job['file_index']}"
            (adir / f"{safe}_raw.json").write_text(
                json.dumps(parsed, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
            )

            empty = detect_empty_items(parsed if isinstance(parsed, dict) else {})
            job["empty_before"] = [e["name"] for e in empty]
            job["fixes"] = []
            job["fix_applied"] = False
            fixed_parsed = parsed
            fix_enabled = bool(cfg.get("fix_empty_answers", True))

            if empty:
                print(f"  🔍 检测到 {len(empty)} 个空答案: {', '.join(e['name'] for e in empty)}", flush=True)
            else:
                print("  ✅ 解析无空答案", flush=True)

            if fix_enabled and empty:
                llm_cfg = resolve_llm_cfg(cfg)
                if not llm_cfg:
                    print(
                        "  ⚠️  需要空题回填但未配置 LLM（secrets.json llm.api_key 或 POLY_LLM_API_KEY），跳过修复",
                        flush=True,
                    )
                    job["fix_skip_reason"] = "no_llm_config"
                else:
                    try:
                        print(f"  🤖 调用 LLM 校验/回填 ({llm_cfg['model']})…", flush=True)
                        t_fix = time.time()
                        fixed_parsed, fixes = fix_empty_answers_with_llm(
                            parsed if isinstance(parsed, dict) else {"content": []},
                            job["path"],
                            llm_cfg,
                            empty_items=empty,
                        )
                        job["fix_s"] = round(time.time() - t_fix, 2)
                        job["fixes"] = fixes
                        job["fix_applied"] = bool(fixes)
                        still = detect_empty_items(fixed_parsed if isinstance(fixed_parsed, dict) else {})
                        job["empty_after"] = [e["name"] for e in still]
                        print(
                            f"  📝 应用 {len(fixes)} 个修正（耗时 {job['fix_s']}s）；"
                            f"仍空 {len(still)} 题"
                            + (f": {', '.join(e['name'] for e in still)}" if still else ""),
                            flush=True,
                        )
                        for fx in fixes:
                            preview = (fx.get("new") or "").replace("\n", " ")[:60]
                            print(f"  ✏️  修正 {fx['name']}: → \"{preview}…\"", flush=True)
                    except Exception as e:
                        job["fix_error"] = str(e)[:500]
                        print(f"  ❌ LLM 回填失败，沿用原始解析: {e}", flush=True)
                        fixed_parsed = parsed
            elif empty and not fix_enabled:
                job["fix_skip_reason"] = "disabled"
                print("  ℹ️  空题回填已关闭（fix_empty_answers=false）", flush=True)

            (adir / f"{safe}_fixed.json").write_text(
                json.dumps(fixed_parsed, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
            )
            job["text_input"] = build_text_input(fixed_parsed if isinstance(fixed_parsed, dict) else {})
            job["analyze_s"] = round(time.time() - t0, 2)
            job["phase"] = "execute"
            print(
                f"  analyze {job['analyze_s']}s text_len={len(job['text_input'] or '')}"
                f" fixes={len(job.get('fixes') or [])}",
                flush=True,
            )
            save_state(cfg, st)
            continue

        if phase == "execute":
            run = cur_run(job)
            t0 = time.time()
            fd = job["file_data"]
            file_list = [
                {
                    "fileId": fd.get("fileId"),
                    "fileName": fd.get("fileName"),
                    "name": fd.get("fileName"),
                    "url": fd.get("ossUrl"),
                    "fileUrl": fd.get("ossUrl"),
                    "ossUrl": fd.get("ossUrl"),
                    "size": fd.get("size"),
                    "suffix": fd.get("suffix"),
                }
            ]
            body = {
                "metadata": {
                    "dimension": "NONE",
                    "instanceNid": instance,
                    "userIds": [user_nid],
                    "version": st["version"],
                    "async": True,
                },
                "sendParams": {
                    "message": {
                        "kind": "message",
                        "parts": [
                            {
                                "kind": "data",
                                "data": {
                                    "fileList": file_list,
                                    "textInput": job.get("text_input"),
                                    "submitType": "FILE_UPLOAD",
                                },
                            }
                        ],
                    }
                },
            }
            r = s.post(f"{HOST}/agents/v1/execute/agent", json=body, timeout=120)
            data = r.json()
            if data.get("code") not in (200, "200"):
                run["error"] = str(data)[:500]
                run["phase"] = "error"
                _advance(st, job, times)
                save_state(cfg, st)
                print(f"  execute FAIL: {data}", flush=True)
                last_code = 0
                continue
            task_id = (data.get("data") or {}).get("id") or (data.get("data") or {}).get("taskId")
            if not task_id:
                run["error"] = f"no task id: {data}"
                run["phase"] = "error"
                _advance(st, job, times)
                save_state(cfg, st)
                continue
            run["task_id"] = task_id
            run["timings"]["execute"] = round(time.time() - t0, 2)
            if ri == 1:
                run["timings"]["upload"] = job.get("upload_s") or 0
                run["timings"]["analyze"] = job.get("analyze_s") or 0
            else:
                run["timings"]["upload"] = 0
                run["timings"]["analyze"] = 0
                run["reused"] = True
            run["poll_t0"] = time.time()
            run["poll_count"] = 0
            run["phase"] = "poll"
            job["phase"] = "poll"
            print(f"  execute {run['timings']['execute']}s task={task_id}", flush=True)
            save_state(cfg, st)
            continue

        if phase == "poll":
            run = cur_run(job)
            task_id = run["task_id"]
            last = None
            done = False
            for _ in range(max_polls):
                if time.time() > deadline:
                    break
                run["poll_count"] = (run.get("poll_count") or 0) + 1
                r = s.post(
                    f"{HOST}/agents/v1/get/task",
                    json={"taskId": task_id, "metadata": {"instanceNid": instance}},
                    timeout=60,
                )
                data = r.json()
                last = data
                d = data.get("data") or {}
                raw_state = (d.get("status") or {}).get("state")
                state = str(raw_state or "")
                arts = d.get("artifacts") or []
                elapsed = time.time() - (run.get("poll_t0") or time.time())
                print(
                    f"  poll#{run['poll_count']} state={raw_state} arts={len(arts)} t={elapsed:.1f}s",
                    flush=True,
                )
                if arts or state in SUCCESS or state.lower() in {x.lower() for x in SUCCESS}:
                    done = True
                    break
                if state in FAIL or state.lower() in {x.lower() for x in FAIL}:
                    done = True
                    break
                time.sleep(interval)

            if not done:
                save_state(cfg, st)
                print(json.dumps({"still_polling": True, "level": level, "run": ri}, ensure_ascii=False))
                last_code = 3
                break

            scores = extract_scores(last or {})
            run["timings"]["poll"] = round(time.time() - (run.get("poll_t0") or time.time()), 2)
            run["total_score"] = scores.get("total")
            run["items"] = scores.get("items") or {}
            run["max_items"] = scores.get("max_items") or {}
            run["state"] = (((last or {}).get("data") or {}).get("status") or {}).get("state")
            run["overall_comment"] = scores.get("overall_comment")
            up = run["timings"].get("upload") or 0
            an = run["timings"].get("analyze") or 0
            ex = run["timings"].get("execute") or 0
            po = run["timings"].get("poll") or 0
            run["timings"]["total"] = round(up + an + ex + po, 2)
            run["phase"] = "done"
            run["completed_at"] = (((last or {}).get("data") or {}).get("status") or {}).get("timestamp")

            (runs_dir(cfg) / f"{level}_run{ri}.json").write_text(
                json.dumps(
                    {"meta": run, "level": level, "filename": job["filename"], "task": last, "scores": scores},
                    ensure_ascii=False,
                    indent=2,
                    default=str,
                ),
                encoding="utf-8",
            )
            print(f"  DONE score={run['total_score']} poll={po}s", flush=True)
            _advance(st, job, times)
            save_state(cfg, st)
            last_code = 0
            continue

        if phase == "done":
            if st["job_index"] + 1 < len(st["jobs"]):
                st["job_index"] += 1
                save_state(cfg, st)
                continue
            st["done"] = True
            save_state(cfg, st)
            break

        print("unknown phase", phase)
        last_code = 1
        break

    print(f"steps={steps} last_code={last_code}", flush=True)
    return last_code


def cmd_status(cfg_path: Path) -> None:
    cfg = load_cfg(cfg_path)
    st = load_state(cfg)
    times = int(st.get("times") or 3)
    done_runs = sum(1 for j in st["jobs"] for r in j["runs"] if r.get("phase") == "done")
    total_runs = len(st["jobs"]) * times
    out = {
        "done": st.get("done"),
        "job_index": st.get("job_index"),
        "progress": f"{done_runs}/{total_runs}",
        "progress_num": done_runs,
        "progress_den": total_runs,
        "wall_so_far": round(time.time() - st["wall_start"], 1),
        "wall_seconds": st.get("wall_seconds"),
        "title": st.get("title"),
        "instance_nid": st.get("instance_nid"),
        "jobs": [
            {
                "level": j["level"],
                "filename": j.get("filename"),
                "job_phase": j["phase"],
                "current_run": j["current_run"],
                "scores": [r.get("total_score") for r in j["runs"]],
                "task_ids": [r.get("task_id") for r in j["runs"]],
                "errors": [r.get("error") for r in j["runs"] if r.get("error")],
            }
            for j in st["jobs"]
        ],
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))


def pop_var(vals: list[float]) -> float | None:
    if not vals:
        return None
    if len(vals) == 1:
        return 0.0
    return statistics.pvariance(vals)


def mean(vals: list[float]) -> float | None:
    if not vals:
        return None
    return sum(vals) / len(vals)


def discover_dims(st: dict) -> list[tuple[str, list[str] | None]]:
    keys: list[str] = []
    for job in st["jobs"]:
        for r in job["runs"]:
            for k in r.get("items") or {}:
                if k not in keys:
                    keys.append(k)
    used = set()
    dims: list[tuple[str, list[str] | None]] = [("总分", None)]

    def take(label, pred):
        group = [k for k in keys if k not in used and pred(k)]
        if group:
            dims.append((label, group))
            used.update(group)

    take("单选题", lambda k: "单选" in k)
    take("多选题", lambda k: "多选" in k)
    take("填空题", lambda k: "填空" in k)
    take("判断题", lambda k: "判断" in k)
    take("问答题", lambda k: "问答" in k)
    take("名词解释", lambda k: "名词" in k or "解释" in k)
    take("论述题", lambda k: "论述" in k)
    for k in keys:
        if k not in used:
            dims.append((k, [k]))
            used.add(k)
    return dims


def dim_score(items: dict, keys: list[str] | None, total) -> float | None:
    if keys is None:
        return float(total) if isinstance(total, (int, float)) else None
    s = 0.0
    any_v = False
    for k in keys:
        v = items.get(k)
        if isinstance(v, (int, float)):
            s += float(v)
            any_v = True
    return s if any_v else None


def cmd_report(cfg_path: Path) -> None:
    cfg = load_cfg(cfg_path)
    st = load_state(cfg)
    dims = discover_dims(st)
    times = int(st.get("times") or 3)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="微软雅黑", color="FFFFFF", bold=True, size=11)
    body_font = Font(name="微软雅黑", size=10)
    bold_font = Font(name="微软雅黑", size=10, bold=True)
    total_fill = PatternFill("solid", fgColor="D6EAF8")
    student_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "评分表"
    headers = ["档次/学生", "评价维度"] + [f"第{i}次" for i in range(1, times + 1)] + ["均值", "方差"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    row = 2
    for job in st["jobs"]:
        student = job.get("filename") or job["level"]
        runs = job["runs"]
        for di, (dim_name, keys) in enumerate(dims):
            scores = []
            for r in runs:
                sc = dim_score(r.get("items") or {}, keys, r.get("total_score"))
                scores.append(sc if isinstance(sc, (int, float)) else None)
            while len(scores) < times:
                scores.append(None)
            valid = [float(x) for x in scores if isinstance(x, (int, float))]
            m = mean(valid)
            v = pop_var(valid)

            ws.cell(row, 1).value = student if di == 0 else None
            ws.cell(row, 2).value = dim_name
            for ci, sc in enumerate(scores[:times], 3):
                ws.cell(row, ci).value = sc
            ws.cell(row, 3 + times).value = round(m, 4) if m is not None else None
            ws.cell(row, 4 + times).value = round(v, 4) if v is not None else None

            for c in range(1, 5 + times):
                cell = ws.cell(row, c)
                cell.border = thin
                cell.font = bold_font if dim_name == "总分" else body_font
                cell.alignment = left if c <= 2 else center
                if dim_name == "总分":
                    cell.fill = total_fill
                if c == 1 and di == 0:
                    cell.fill = student_fill
                    cell.font = bold_font
            row += 1
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    for i in range(3, 5 + times):
        ws.column_dimensions[get_column_letter(i)].width = 10
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "C2"

    ws2 = wb.create_sheet("批阅明细")
    item_keys = []
    for job in st["jobs"]:
        for r in job["runs"]:
            for k in r.get("items") or {}:
                if k not in item_keys:
                    item_keys.append(k)
    h2 = ["档次", "轮次", "总分", "状态", "task_id"] + item_keys
    for c, h in enumerate(h2, 1):
        cell = ws2.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    ri = 2
    for job in st["jobs"]:
        for r in job["runs"]:
            vals = [job["level"], r.get("run_index"), r.get("total_score"), r.get("state"), r.get("task_id")]
            for k in item_keys:
                vals.append((r.get("items") or {}).get(k))
            for c, v in enumerate(vals, 1):
                cell = ws2.cell(ri, c, v)
                cell.border = thin
                cell.alignment = center
                cell.font = body_font
            ri += 1
    for c in range(1, len(h2) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 12 if c != 5 else 38

    ws3 = wb.create_sheet("说明")
    fix_lines = []
    for job in st["jobs"]:
        eb = job.get("empty_before") or []
        ea = job.get("empty_after")
        nf = len(job.get("fixes") or [])
        if eb or nf:
            fix_lines.append(
                f"  - {job.get('filename') or job['level']}: 空题前={len(eb)} 修正={nf}"
                + (f" 空题后={len(ea)}" if ea is not None else "")
                + (f" skip={job.get('fix_skip_reason')}" if job.get("fix_skip_reason") else "")
                + (f" err={job.get('fix_error')}" if job.get("fix_error") else "")
            )
    notes = [
        f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"标题: {st.get('title') or '-'}",
        f"Instance: {st.get('instance_nid')}",
        f"Course: {st.get('course_id') or '-'}",
        f"每份批阅次数: {times}",
        "策略: 每份 upload×1 + analyze×1 (+ 可选 LLM 空题回填) + execute×N",
        f"空题回填: fix_empty_answers={cfg.get('fix_empty_answers', True)}",
        "方差: 总体方差（÷n）",
        f"完成: done={st.get('done')} wall≈{st.get('wall_seconds')}s",
        "解析修复明细:",
        *(fix_lines or ["  (无空题或未触发修复)"]),
    ]
    for i, t in enumerate(notes, 1):
        ws3.cell(i, 1, t).font = body_font
    ws3.column_dimensions["A"].width = 100

    excel_name = cfg.get("excel_name") or "评分表.xlsx"
    excel = Path(cfg["out_dir"]) / excel_name
    wb.save(excel)

    summary = []
    for job in st["jobs"]:
        totals = [r.get("total_score") for r in job["runs"] if isinstance(r.get("total_score"), (int, float))]
        summary.append(
            {
                "level": job["level"],
                "filename": job.get("filename"),
                "scores": [r.get("total_score") for r in job["runs"]],
                "mean": round(mean(totals), 4) if totals else None,
                "var": round(pop_var(totals), 4) if totals else None,
            }
        )
    summary_path = Path(cfg["out_dir"]) / "summary.json"
    summary_path.write_text(
        json.dumps(
            {
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "files": summary,
                "wall_seconds": st.get("wall_seconds"),
                "instance_nid": st.get("instance_nid"),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"excel": str(excel), "summary": str(summary_path), "files": summary}, ensure_ascii=False, indent=2))


def cmd_run(cfg_path: Path, max_loops: int = 200) -> None:
    for i in range(max_loops):
        cfg = load_cfg(cfg_path)
        st = load_state(cfg)
        if st.get("done"):
            break
        code = cmd_next(cfg_path)
        if code not in (0, 3):
            sys.exit(code)
        time.sleep(0.2)
    cmd_report(cfg_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("command", choices=["init", "next", "status", "report", "run"])
    ap.add_argument("--config", required=True, help="job config json path")
    args = ap.parse_args()
    cfg_path = Path(args.config)
    if args.command == "init":
        cmd_init(cfg_path)
    elif args.command == "next":
        code = cmd_next(cfg_path)
        sys.exit(0 if code in (0, 3) else code)
    elif args.command == "status":
        cmd_status(cfg_path)
    elif args.command == "report":
        cmd_report(cfg_path)
    elif args.command == "run":
        cmd_run(cfg_path)


if __name__ == "__main__":
    main()
