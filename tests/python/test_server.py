# server.py 测试：parse_file_content 单测 + /api/start-harness 端点集成测试
# 依赖（已在受管 venv 安装）：fastapi, python-multipart, python-docx, pypdf
#
# 注意：server.py 当前为异步任务模型——POST 立即返回 {run_id, status:"pending"}，
# 后台线程跑 harness，结果通过 GET /api/runs/{run_id} 轮询。测试据此对齐。
import io
import time

from fastapi.testclient import TestClient

from server import app, parse_file_content


# ---------- parse_file_content ----------

def test_parse_plain_text():
    out = parse_file_content("note.txt", "hello 世界".encode())
    assert out == "hello 世界"


def test_parse_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("段落一内容")
    doc.add_paragraph("段落二内容")
    buf = io.BytesIO()
    doc.save(buf)
    out = parse_file_content("plan.docx", buf.getvalue())
    assert "段落一内容" in out
    assert "段落二内容" in out


def test_parse_pdf():
    out = parse_file_content("doc.pdf", _make_minimal_pdf("HelloPDFExtract"))
    assert "HelloPDFExtract" in out


def _make_minimal_pdf(text: str) -> bytes:
    """构造一个带正确 xref 偏移的最小合法 PDF（离线、无第三方依赖）。"""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
    ]
    stream = f"BT /F1 12 Tf 50 150 Td ({text}) Tj ET".encode("latin-1")
    stream_obj = (
        b"<< /Length " + str(len(stream)).encode()
        + b" >>\nstream\n" + stream + b"\nendstream"
    )
    objects.append(stream_obj)
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\n"
    out += b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF\n"
    return out


# ---------- /api/start-harness 集成测试（异步任务模型） ----------

client = TestClient(app)


def _wait_for_run(run_id: str, timeout: float = 15.0) -> dict:
    """轮询 GET /api/runs/{run_id}，直到 completed/failed。"""
    deadline = time.time() + timeout
    while time.time() < deadline:
        r = client.get(f"/api/runs/{run_id}")
        assert r.status_code == 200
        job = r.json()
        if job["status"] in ("completed", "failed"):
            return job
        time.sleep(0.05)
    raise AssertionError(f"run {run_id} 未在 {timeout}s 内结束")


def test_root_serves_html():
    r = client.get("/")
    assert r.status_code == 200
    assert "text/html" in r.headers["content-type"]


def test_start_harness_returns_run_id_immediately():
    r = client.post("/api/start-harness", data={"text": "x"})
    assert r.status_code == 200
    body = r.json()
    assert "run_id" in body and isinstance(body["run_id"], str) and body["run_id"]
    assert body["status"] == "pending"


def test_start_harness_threshold_clamped_in_result():
    # 越界 threshold 应在 HermesAgent.run 内被夹到 [1,100]
    r = client.post("/api/start-harness", data={"text": "x", "threshold": 999})
    run_id = r.json()["run_id"]
    job = _wait_for_run(run_id)
    assert job["status"] == "completed"
    assert 1 <= job["result"]["threshold"] <= 100


def test_start_harness_with_uploaded_txt_completes():
    r = client.post(
        "/api/start-harness",
        files={"document": ("task.txt", b"uploaded training doc content")},
        data={"text": ""},
    )
    run_id = r.json()["run_id"]
    job = _wait_for_run(run_id)
    assert job["status"] == "completed"
    # task_summary 应来自上传文档内容（经 parse_file_content）
    assert job["result"]["task_summary"]


def test_get_unknown_run_returns_404():
    r = client.get("/api/runs/does-not-exist")
    assert r.status_code == 404
